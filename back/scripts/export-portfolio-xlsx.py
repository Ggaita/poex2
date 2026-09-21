import openpyxl, json, re, pathlib
path = r"C:\Users\gaita\Downloads\Base_de_Datos_Cartera_de_Inversiones_y_Parques_Industriales.xlsx"
out = pathlib.Path(r"D:\2.PROGRAMACIÓN\5. FIERROTECH\1. WEB\1. POEX\poex2\back\data")
out.mkdir(parents=True, exist_ok=True)
wb = openpyxl.load_workbook(path, data_only=True)

def clean(v):
    if v is None:
        return None
    if isinstance(v, float) and v == int(v):
        return int(v)
    if isinstance(v, str):
        t = v.strip()
        return t or None
    return v

def slugify(s):
    s = str(s or "").lower()
    repl = {"á":"a","é":"e","í":"i","ó":"o","ú":"u","ñ":"n","ü":"u"}
    for a,b in repl.items():
        s = s.replace(a,b)
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s[:80] or "item"

s = wb["Oportunidades de Inversión"]
rows = list(s.iter_rows(values_only=True))
h = [str(c).strip() if c is not None else "" for c in rows[0]]
ops = []
for i, r in enumerate(rows[1:], start=1):
    if not any(c is not None and str(c).strip() for c in r):
        continue
    d = {h[j]: clean(r[j]) if j < len(r) else None for j in range(len(h))}
    ext_id = str(d.get("ID") or f"X{i}")
    title = d.get("Nombre / Título") or f"Oportunidad {ext_id}"
    desc = d.get("Descripción") or title
    ops.append({
        "externalId": ext_id,
        "excelType": d.get("Tipo de Oportunidad"),
        "title": title,
        "description": desc,
        "sector": d.get("Cadena / Sector") or "General",
        "locality": d.get("Localidad / Región") or "Chaco",
        "nature": d.get("Naturaleza (Pública/Privada/Mixta)"),
        "maturity": d.get("Nivel de Madurez"),
        "activation": d.get("Mecanismo de Activación Identificado"),
        "source": d.get("Fuente / Entrevista"),
        "observations": d.get("Observaciones"),
        "slug": f"{slugify(ext_id)}-{slugify(title)}"
    })

s2 = wb["Parques Industriales"]
rows2 = list(s2.iter_rows(values_only=True))
h2 = [str(c).strip() if c is not None else "" for c in rows2[0]]
parks = []
for i, r in enumerate(rows2[1:], start=1):
    if not any(c is not None and str(c).strip() for c in r):
        continue
    d = {h2[j]: clean(r[j]) if j < len(r) else None for j in range(len(h2))}
    name = d.get("Nombre")
    if not name:
        continue
    parks.append({
        "number": d.get("N°") or i,
        "name": name,
        "administration": d.get("Administración"),
        "progressStatus": d.get("Estado de Avance") or "Sin dato",
        "locality": d.get("Localidad") or "Chaco",
        "department": d.get("Departamento"),
        "yearCreated": d.get("Año de Creación"),
        "surfaceHa": d.get("Superficie (Has)"),
        "measuredPlots": d.get("Parcelas Mensuradas"),
        "settledCompanies": d.get("Empresas Radicadas"),
        "infrastructure": d.get("Obras de Infraestructura"),
        "renpiStatus": d.get("Estado ReNPI"),
        "observations": d.get("Observaciones"),
        "slug": slugify(name)
    })

(out / "investment-opportunities-seed.json").write_text(json.dumps(ops, ensure_ascii=False, indent=2), encoding="utf-8")
(out / "industrial-parks-seed.json").write_text(json.dumps(parks, ensure_ascii=False, indent=2), encoding="utf-8")
print("ops", len(ops), "parks", len(parks))
